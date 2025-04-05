from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# 创建一个新的 Word 文档
doc = Document()

# 设置标题格式和样式
def add_title(text, level=1):
    """
    添加标题。
    :param text: 标题文本
    :param level: 标题级别
    """
    doc.add_heading(text, level=level)

# 设置正文段落格式
def add_paragraph(text, indent_level=0):
    """
    添加正文段落。
    :param text: 段落文本
    :param indent_level: 缩进级别
    """
    para = doc.add_paragraph(text)
    para.style.font.name = "宋体"
    para.style.font.size = Pt(12)
    para.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    if indent_level > 0:
        para.paragraph_format.left_indent = Pt(indent_level * 14.4)  # 缩进一阶对应0.5厘米

# 添加论文大纲内容
add_title("第一章 绪论", level=1)
add_title("1.1 研究背景与意义", level=2)
add_paragraph("传统舆情分析在信息检索和数据处理方面的不足", indent_level=1)
add_paragraph("大语言模型（Large Language Models, LLM）在舆情分析中的潜在价值", indent_level=1)
add_paragraph("面向舆情分析的领域优化与场景需求的重要性", indent_level=1)

add_title("1.2 国内外研究现状", level=2)
add_paragraph("大语言模型在信息检索与舆情分析领域的应用进展", indent_level=1)
add_paragraph("现有舆情分析系统的技术路线与局限性", indent_level=1)
add_paragraph("信息冗余、噪声与实时性问题的研究现状", indent_level=1)

add_title("1.3 研究目标与内容", level=2)
add_paragraph("提出一个基于大模型的舆情事件分析系统的总体目标", indent_level=1)
add_paragraph("解决舆情数据采集、查询重写、文本总结、可视化展示等关键问题", indent_level=1)
add_paragraph("研究内容概述：", indent_level=1)
add_paragraph("检索与数据处理方法", indent_level=2)
add_paragraph("舆情事件多维度分析与生成", indent_level=2)
add_paragraph("系统设计与实现", indent_level=2)
add_paragraph("实验评估与性能分析", indent_level=2)

add_title("1.4 论文结构安排", level=2)
add_paragraph("第二章：相关理论与技术概述", indent_level=1)
add_paragraph("第三章：基于大模型的文本检索方法研究", indent_level=1)
add_paragraph("第四章：基于大模型的舆情分析方法研究", indent_level=1)
add_paragraph("第五章：舆情事件分析系统的设计与实现", indent_level=1)
add_paragraph("第六章：总结与展望", indent_level=1)

add_title("第二章 相关理论与技术概述", level=1)
add_title("2.1 大语言模型理论基础", level=2)
add_paragraph("自然语言处理领域的范式转变", indent_level=1)
add_paragraph("大语言模型对任务迁移学习的推动作用", indent_level=1)
add_paragraph("提示工程（Prompt Engineering）的概念与演进", indent_level=1)

add_title("2.2 信息检索系统的技术演进", level=2)
add_paragraph("传统检索模型（TF-IDF、BM25）与深度学习检索模型（Siamese Networks、Transformer）的对比", indent_level=1)
add_paragraph("语义匹配与语义索引在舆情分析领域的应用", indent_level=1)
add_paragraph("多搜索引擎融合与检索多样性保障", indent_level=1)

add_title("2.3 舆情分析相关概念", level=2)
add_paragraph("舆情分析定义与核心任务（事件识别、观点挖掘、影响评估等）", indent_level=1)
add_paragraph("舆情分析常见的挑战：信息实时性、噪声过滤、领域歧义等", indent_level=1)
add_paragraph("基于 LLM 的舆情分析价值：多模态信息整合、深度语义理解", indent_level=1)

add_title("2.4 数据获取与处理技术", level=2)
add_paragraph("爬虫技术与多搜索引擎接口调用", indent_level=1)
add_paragraph("数据清洗、去重与特征提取", indent_level=1)
add_paragraph("向量化表示与向量数据库（如 FAISS、Milvus、InMemoryVectorStore）", indent_level=1)

add_title("第三章 基于大模型的文本检索方法研究", level=1)
add_title("3.1 问题描述与研究目标", level=2)
add_paragraph("检索质量优化的主要挑战：查询歧义、冗余结果、竞价排名带来的噪声", indent_level=1)
add_paragraph("设计有效的检索流程：从用户意图识别到多引擎结果融合", indent_level=1)

add_title("3.2 检索优化方案设计", level=2)
add_title("3.2.1 提升用户查询质量的方法", level=3)
add_paragraph("用户意图识别与查询重写", indent_level=1)
add_paragraph("提示工程在查询改写中的应用：链式思维、分段提示等", indent_level=1)
add_paragraph("减少歧义的策略：领域术语扩展、子查询生成", indent_level=1)

add_title("3.2.2 提升结果排序的相关性", level=3)
add_paragraph("多搜索引擎检索结果的整合与去重", indent_level=1)
add_paragraph("相关性重排序：基于大模型的重排序与传统方法结合", indent_level=1)
add_paragraph("语义向量检索与文本片段拼接", indent_level=1)

add_title("3.3 实验设计与效果评估", level=2)
add_title("3.3.1 查询改写实验与结果分析", level=3)
add_paragraph("查询改写前后检索效果对比指标（MRR、nDCG 等）", indent_level=1)

# 补充后续章节内容并继续完善文档

# 第四章
add_title("第四章 基于大模型的舆情分析方法研究", level=1)
add_title("4.1 问题描述与研究目标", level=2)
add_paragraph("舆情分析在信息抽取、观点识别、影响力分析等方面的难点", indent_level=1)
add_paragraph("结合提示工程设计多轮分析策略", indent_level=1)

add_title("4.2 方法设计与实现", level=2)
add_title("4.2.1 舆情事件分析任务分解方法", level=3)
add_paragraph("事件概述、时间线、舆情态度及潜在影响的分步骤分析", indent_level=1)
add_paragraph("链式提示如何引导大模型聚焦不同细分任务", indent_level=1)

add_title("4.2.2 多轮交互优化分析流程", level=3)
add_paragraph("分层提示设计：初始分析、深度挖掘、总结生成", indent_level=1)
add_paragraph("知识扩展与子问题挖掘：使用长上下文模型（如 Qwen-2.5 128k）进行文本清洗", indent_level=1)

add_title("4.3 实验设计与评估", level=2)
add_title("4.3.1 分析准确性对比实验", level=3)
add_paragraph("基准模型（GPT-3.5/4 等）在舆情分析场景下的表现", indent_level=1)
add_paragraph("分步分析 vs. 一次性生成的准确度对比", indent_level=1)

add_title("4.3.2 多模型协同分析的效果验证", level=3)
add_paragraph("长上下文模型 + 高性能模型的协同流程", indent_level=1)
add_paragraph("不同组合方案在舆情分析准确度与效率方面的对比", indent_level=1)

add_title("4.3.3 舆情领域具体案例研究", level=3)
add_paragraph("公共卫生事件、国际冲突、企业危机等典型舆情案例", indent_level=1)
add_paragraph("分析结果与人工标注或官方数据之间的对比", indent_level=1)

# 第五章
add_title("第五章 基于大模型的舆情事件分析系统设计与实现", level=1)
add_title("5.1 系统需求与业务分析", level=2)
add_paragraph("针对政府、企业、研究机构的核心需求", indent_level=1)
add_paragraph("实时性、高并发、多用户管理与系统鲁棒性", indent_level=1)

add_title("5.2 系统设计", level=2)
add_title("5.2.1 系统整体架构", level=3)
add_paragraph("前端交互层、中间层处理逻辑、后端数据层", indent_level=1)
add_paragraph("WebSocketManager 模块管理实时任务与通信", indent_level=1)

add_title("5.2.2 关键模块功能与交互设计", level=3)
add_paragraph("检索模块：多搜索引擎统一接口与重排序机制", indent_level=1)
add_paragraph("数据处理模块：爬虫、清洗、向量化、索引构建", indent_level=1)
add_paragraph("生成式对话模块：LLM", indent_level=1)

# 第六章
add_title("第六章 总结与展望", level=1)
add_title("6.1 工作总结", level=2)
add_paragraph("总结本文提出的基于大语言模型的舆情事件分析方法及系统实现", indent_level=1)
add_paragraph("验证方法在检索、分析、生成等环节的有效性", indent_level=1)

add_title("6.2 未来展望", level=2)
add_paragraph("探讨未来研究中对多模态舆情分析、实时更新与模型性能优化的可能性", indent_level=1)
add_paragraph("提出改进方向，如更高效的数据处理、更智能的交互设计等", indent_level=1)


# 更多内容可按需要添加...

# 保存为Word文档
file_path = "/Users/lgq/Desktop/毕设文件/毕设论文资料/基于大语言模型的舆情分析论文大纲.docx"
doc.save(file_path)
